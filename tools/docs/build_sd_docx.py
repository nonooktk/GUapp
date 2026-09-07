# -*- coding: utf-8 -*-
# 設計仕様書 P1 本編（Markdown 正本）→ docx。Mermaid は事前レンダリング済み PNG を埋め込む
import re, os, sys
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.stdout.reconfigure(encoding='utf-8')

MD = r'C:\Users\m-oya\Desktop\GUapp\docs\03_設計仕様書\GU_ECsite_設計仕様書_P1_draft-v2.md'
IMG_DIR = r'C:\Users\m-oya\AppData\Local\Temp\claude\C--Users-m-oya-Desktop-DocsMaker\03a8bbb3-8c45-4cd2-ac2f-5c0734de2a32\scratchpad\mmd'
OUT = os.environ.get('SD_OUT', r'C:\Users\m-oya\Desktop\DocsMaker\成果物\GU_ECsite_設計仕様書_P1_draft-v2_20260907.docx')
SKIP_IMG = os.environ.get('SKIP_IMG') == '1'
SKIP_CODE = os.environ.get('SKIP_CODE') == '1'
FONT = 'メイリオ'
MONO = 'ＭＳ ゴシック'
HEADER_FILL = 'F2F2F2'

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Mm(297), Mm(210)
sec.top_margin = sec.bottom_margin = Mm(25.4)
sec.left_margin = sec.right_margin = Mm(19.05)
sec.header_distance = Mm(15); sec.footer_distance = Mm(17.5)
sec.different_first_page_header_footer = True
TEXT_W = sec.page_width - sec.left_margin - sec.right_margin

# ---- Word の表示設定：常に印刷レイアウトで開く・互換モードを解除（Word 2013+）----
_settings = doc.settings.element
for _old in _settings.findall(qn('w:view')):
    _settings.remove(_old)
_view = OxmlElement('w:view'); _view.set(qn('w:val'), 'print')
_settings.insert(0, _view)
_compat = _settings.find(qn('w:compat'))
if _compat is None:
    _compat = OxmlElement('w:compat'); _settings.append(_compat)
for _cs in list(_compat.findall(qn('w:compatSetting'))):
    if _cs.get(qn('w:name')) == 'compatibilityMode':
        _compat.remove(_cs)
_cs = OxmlElement('w:compatSetting')
_cs.set(qn('w:name'), 'compatibilityMode'); _cs.set(qn('w:uri'), 'http://schemas.microsoft.com/office/word'); _cs.set(qn('w:val'), '15')
_compat.append(_cs)

def set_ea(rPr, name):
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name); rFonts.set(qn('w:hAnsi'), name); rFonts.set(qn('w:eastAsia'), name)

def style_font(style, size, bold=False, color='000000'):
    style.font.name = FONT; style.font.size = Pt(size); style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    set_ea(style.element.get_or_add_rPr(), FONT)

normal = doc.styles['Normal']; style_font(normal, 10.5)
normal.paragraph_format.line_spacing = 1.1; normal.paragraph_format.space_after = Pt(4)
h1 = doc.styles['Heading 1']; style_font(h1, 14, True)
h1.paragraph_format.space_before = Pt(16); h1.paragraph_format.space_after = Pt(8)
h2 = doc.styles['Heading 2']; style_font(h2, 12, True)
h2.paragraph_format.space_before = Pt(12); h2.paragraph_format.space_after = Pt(6)
h3 = doc.styles['Heading 3']; style_font(h3, 11, True)
h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(4)

hdr = sec.header.paragraphs[0]; hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hdr.add_run('GU ECサイト 設計仕様書（P1 本編）'); r.font.name = FONT; r.font.size = Pt(8); r.font.italic = True
r.font.color.rgb = RGBColor.from_string('666666'); set_ea(r._element.get_or_add_rPr(), FONT)
ftr = sec.footer.paragraphs[0]; ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
ftr.add_run().font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
_fr = OxmlElement('w:r'); _ft = OxmlElement('w:t'); _ft.text = '1'; _fr.append(_ft); fld.append(_fr); ftr._p.append(fld)

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')

def add_runs(p, text, size=10.5, color=None, base_bold=False):
    for part in INLINE.split(text):
        if not part: continue
        bold = base_bold; mono = False
        if part.startswith('**') and part.endswith('**'):
            part = part[2:-2]; bold = True
        elif part.startswith('`') and part.endswith('`'):
            part = part[1:-1]; mono = True
        r = p.add_run(part)
        r.font.name = 'Consolas' if mono else FONT
        r.font.size = Pt(size); r.font.bold = bold
        if color: r.font.color.rgb = RGBColor.from_string(color)
        set_ea(r._element.get_or_add_rPr(), MONO if mono else FONT)

def para(text, size=10.5, bold=False, align=None, space_after=4, color=None, indent=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None: p.paragraph_format.left_indent = Pt(indent)
    add_runs(p, text, size=size, color=color, base_bold=bold)
    return p

def code_block(lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Pt(10)
        r = p.add_run(ln if ln else ' ')
        r.font.name = MONO; r.font.size = Pt(7.5)
        set_ea(r._element.get_or_add_rPr(), MONO)
    para('', size=4, space_after=2)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def cell_text(cell, text, size, bold=False, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1); p.paragraph_format.line_spacing = 1.05
    if align is not None: p.alignment = align
    add_runs(p, text, size=size, base_bold=bold)

def visual_len(s):
    s = re.sub(r'\*\*|`', '', s)
    return sum(2 if ord(c) > 0x2E7F else 1 for c in s)

def add_table(rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    body = [r + [''] * (ncol - len(r)) if len(r) < ncol else r[:ncol] for r in body]
    size = 9.5 if ncol <= 3 else (8.5 if ncol <= 5 else 8)
    # 列幅：各列の代表長（ヘッダーと本文の 90 パーセンタイル）に比例、最小幅を確保
    lens = []
    for j in range(ncol):
        col = sorted(visual_len(r[j]) for r in [header] + body)
        rep = col[int(len(col) * 0.9)] if col else 4
        lens.append(max(6, min(rep, 60)))
    total = sum(lens)
    widths = [max(int(TEXT_W * l / total), int(TEXT_W * 0.07)) for l in lens]
    scale = TEXT_W / sum(widths); widths = [int(w * scale) for w in widths]
    tbl = doc.add_table(rows=1 + len(body), cols=ncol)
    tbl.style = doc.styles['Table Grid']; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofit = False
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top', 15), ('bottom', 15), ('start', 50), ('end', 50)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(w)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    _look = tbl._tbl.tblPr.find(qn('w:tblLook'))
    (_look.addprevious(mar) if _look is not None else tbl._tbl.tblPr.append(mar))
    for j, h in enumerate(header):
        c = tbl.rows[0].cells[j]; c.width = widths[j]
        cell_text(c, h, size, bold=True); shade(c, HEADER_FILL)
    for i, row in enumerate(body, start=1):
        for j, v in enumerate(row):
            c = tbl.rows[i].cells[j]; c.width = widths[j]
            cell_text(c, v, size)
    para('', size=6, space_after=0)

def add_image(idx):
    path = os.path.join(IMG_DIR, f'diagram_{idx}.png')
    if not os.path.exists(path):
        para(f'（図 {idx} は未生成）', color='C00000'); return
    # ページ高さを超える縦長図は高さ 225mm に収める（超えると Word がレイアウトできず下書き表示に落ちる）
    from PIL import Image
    with Image.open(path) as im:
        w_px, h_px = im.size
    max_h = Mm(225)
    width = TEXT_W
    if TEXT_W * h_px / w_px > max_h:
        width = int(max_h * w_px / h_px)
    doc.add_picture(path, width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para('', size=6, space_after=0)

def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]

# ---------- パース ----------
lines = open(MD, encoding='utf-8').read().splitlines()
i = 0; mermaid_idx = 0
meta = {}
# 先頭のタイトルとメタ表を表紙に使う
title = lines[0].lstrip('# ').strip()
i = 1
while i < len(lines) and not lines[i].startswith('|'): i += 1
while i < len(lines) and lines[i].startswith('|'):
    cells = split_row(lines[i])
    if len(cells) >= 2 and not set(cells[0]) <= set('-: '):
        if cells[0] not in ('項目',): meta[cells[0]] = cells[1]
    i += 1

# ---------- 表紙 ----------
for _ in range(7): para('', space_after=0)
para('GU ECサイト', size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para('設計仕様書（P1 本編）', size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para(f"版 {meta.get('版','')}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
for _ in range(3): para('', space_after=0)
info = doc.add_table(rows=len(meta), cols=2); info.style = doc.styles['Table Grid']
info.alignment = WD_TABLE_ALIGNMENT.CENTER; info.autofit = False
for k_i, (k, v) in enumerate(meta.items()):
    c0, c1 = info.rows[k_i].cells
    c0.width = int(TEXT_W * 0.22); c1.width = int(TEXT_W * 0.60)
    cell_text(c0, k, 10, align=WD_ALIGN_PARAGRAPH.CENTER); cell_text(c1, v, 10)
doc.add_page_break()

# ---------- 本文 ----------
table_buf = []
def flush_table():
    global table_buf
    if table_buf:
        rows = [split_row(l) for l in table_buf if not set(l.replace('|', '').strip()) <= set('-: ')]
        if rows: add_table(rows)
        table_buf = []

n = len(lines)
while i < n:
    ln = lines[i]
    if ln.startswith('```'):
        flush_table()
        lang = ln[3:].strip()
        j = i + 1; block = []
        while j < n and not lines[j].startswith('```'):
            block.append(lines[j]); j += 1
        if lang == 'mermaid':
            mermaid_idx += 1
            if SKIP_IMG: para(f'[図 {mermaid_idx}]')
            else: add_image(mermaid_idx)
        else:
            if SKIP_CODE: para('[code block]')
            else: code_block(block)
        i = j + 1; continue
    if ln.startswith('|'):
        table_buf.append(ln); i += 1; continue
    flush_table()
    s = ln.strip()
    if not s:
        i += 1; continue
    if ln.startswith('#### '):
        doc.add_paragraph(ln[5:].strip(), style='Heading 3')
    elif ln.startswith('### '):
        doc.add_paragraph(ln[4:].strip(), style='Heading 2')
    elif ln.startswith('## '):
        doc.add_paragraph(ln[3:].strip(), style='Heading 1')
    elif re.match(r'^\s*- ', ln):
        para('・ ' + re.sub(r'^\s*- ', '', ln), space_after=3, indent=6)
    elif re.match(r'^\d+\. ', s):
        para(s, space_after=3, indent=6)
    elif s.startswith('※'):
        para(s, size=9, color='595959', space_after=3, indent=14)
    else:
        para(s)
    i += 1
flush_table()

doc.save(OUT)
print('saved:', OUT)
d2 = Document(OUT)
print('tables:', len(d2.tables), '/ images:', len(d2.inline_shapes), '/ mermaid blocks:', mermaid_idx)
